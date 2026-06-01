import asyncio
import openpyxl
from openpyxl.styles import Font
import aiosqlite
from datetime import datetime, date
from collections import defaultdict

DB_NAME = "workers_time.db"


def _build_excel_user(rows, filename: str):
    """Создает Excel-файл для конкретного сотрудника (один лист)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Мой учёт времени"

    headers = ["ID", "ФИО", "Дата", "Приход", "Уход", "Отработано"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total_seconds = 0

    for row in rows:
        uid, name, w_date, arr, dep = row
        hours_worked = "Не завершено"

        if arr and dep:
            t1 = datetime.strptime(arr, "%H:%M:%S")
            t2 = datetime.strptime(dep, "%H:%M:%S")
            delta = t2 - t1
            total_seconds += delta.total_seconds()

            hrs = int(delta.total_seconds() // 3600)
            mins = int((delta.total_seconds() % 3600) // 60)
            hours_worked = f"{hrs}ч {mins}м"

        ws.append([uid, name, w_date, arr or "-", dep or "-", hours_worked])

    # Добавляем итоговую строку для сотрудника
    total_hrs = int(total_seconds // 3600)
    total_mins = int((total_seconds % 3600) // 60)
    ws.append([])
    ws.append(["ИТОГО ЗА МЕСЯЦ:", "", "", "", "", f"{total_hrs}ч {total_mins}м"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=6).font = Font(bold=True)

    wb.save(filename)


def _build_excel_admin(rows, filename: str):
    """Создает Excel-файл для администратора с двумя листами (Итоги + Логи)."""
    wb = openpyxl.Workbook()

    # --- ЛИСТ 1: ОБЩАЯ СТАТИСТИКА ЗА МЕСЯЦ ---
    ws_summary = wb.active
    ws_summary.title = "Общие итоги"
    ws_summary.append(["ID сотрудника", "ФИО сотрудника", "Всего отработано за месяц"])
    ws_summary.cell(row=1, column=1).font = Font(bold=True)
    ws_summary.cell(row=1, column=2).font = Font(bold=True)
    ws_summary.cell(row=1, column=3).font = Font(bold=True)

    # --- ЛИСТ 2: ПОДРОБНЫЕ ЛОГИ ВСЕХ СМЕН ---
    ws_logs = wb.create_sheet(title="Подробные логи")
    headers = ["ID", "ФИО", "Дата", "Приход", "Уход", "Отработано"]
    ws_logs.append(headers)
    for cell in ws_logs[1]:
        cell.font = Font(bold=True)

    # Словарь для подсчета суммарного времени по каждому человеку
    # Структура: { user_id: { "name": "ФИО", "seconds": 0 } }
    user_totals = defaultdict(lambda: {"name": "", "seconds": 0})

    # Заполняем подробные логи и одновременно ведем математический подсчет итогов
    for row in rows:
        uid, name, w_date, arr, dep = row
        hours_worked = "Не завершено"

        if arr and dep:
            t1 = datetime.strptime(arr, "%H:%M:%S")
            t2 = datetime.strptime(dep, "%H:%M:%S")
            delta = t2 - t1

            # Плюсуем секунды конкретному сотруднику в общую копилку
            user_totals[uid]["name"] = name
            user_totals[uid]["seconds"] += delta.total_seconds()

            hrs = int(delta.total_seconds() // 3600)
            mins = int((delta.total_seconds() % 3600) // 60)
            hours_worked = f"{hrs}ч {mins}м"
        else:
            if name:
                user_totals[uid]["name"] = name

        ws_logs.append([uid, name, w_date, arr or "-", dep or "-", hours_worked])

    # Заполняем первый лист ("Общие итоги") посчитанными агрегированными данными
    for uid, data in user_totals.items():
        total_sec = data["seconds"]
        t_hrs = int(total_sec // 3600)
        t_mins = int((total_sec % 3600) // 60)
        ws_summary.append([uid, data["name"], f"{t_hrs}ч {t_mins}м"])

    wb.save(filename)


async def generate_monthly_report(for_admin: bool = False, user_id: int = None) -> str:
    """Генерирует Excel-файл. Перенаправляет логику в зависимости от роли."""
    now = datetime.now()
    start_date = date(now.year, now.month, 1).isoformat()
    end_date = (
        date(now.year, now.month + 1, 1).isoformat()
        if now.month < 12
        else date(now.year + 1, 1, 1).isoformat()
    )

    async with aiosqlite.connect(DB_NAME) as db:
        if for_admin:
            query = """
                SELECT u.user_id, u.full_name, t.work_date, t.arrival_time, t.departure_time
                FROM time_logs t JOIN users u ON t.user_id = u.user_id
                WHERE t.work_date >= ? AND t.work_date < ?
                ORDER BY u.full_name ASC, t.work_date ASC
            """
            params = (start_date, end_date)
            filename = f"report_admin_{now.strftime('%m_%Y')}.xlsx"
        else:
            query = """
                SELECT u.user_id, u.full_name, t.work_date, t.arrival_time, t.departure_time
                FROM time_logs t JOIN users u ON t.user_id = u.user_id
                WHERE t.user_id = ? AND t.work_date >= ? AND t.work_date < ?
                ORDER BY t.work_date ASC
            """
            params = (user_id, start_date, end_date)
            filename = f"report_{user_id}_{now.strftime('%m_%Y')}.xlsx"

        async with db.execute(query, params) as cursor:
            rows = await cursor.fetchall()
            if not rows:
                return None

    # Перенаправляем сборку в разные функции потоков процессора
    if for_admin:
        await asyncio.to_thread(_build_excel_admin, rows, filename)
    else:
        await asyncio.to_thread(_build_excel_user, rows, filename)

    return filename
